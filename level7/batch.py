"""Batch marker orders: a list (Excel .xlsx or CSV) of markers, each with its sizes from one or more design
files; LEVEL7 builds, nests, saves and exports them one after the other.

List layout - one row per (marker, file, size); marker-level settings are taken from the first row of the
marker that has them:

    marker | file       | size | qty | width | material | formats  | time | spacing | engine
    (optional: buffer_x | buffer_y | shrink_x | shrink_y — extra buffer between pieces in mm, shrinkage in %)
    M001   | pants.mdl  | 44   | 2   | 1500  | O        | plt;iso  | 120  | 0       | auto
    M001   | pants.mdl  | 46   | 1
    M001   | shirt.mdl  | M    | 1
    M002   | shirt.mdl  | L    | 3   | 1440  | O        | dxf      | 60

Header names may also be Arabic (ماركر, ملف, مقاس, كمية, عرض, خامة, صيغ, وقت, مسافة,
بفر طول, بفر عرض, انكماش طول, انكماش عرض). File paths may be
relative to the list's folder. Formats: plt, dxf, iso (separated by ; , or space).
"""
from __future__ import annotations

import csv
import os
import re
import time
from dataclasses import dataclass, field
from pathlib import Path

HEADERS = {
    "marker": ("marker", "ماركر", "الماركر", "name"),
    "file": ("file", "ملف", "الملف", "model", "موديل"),
    "size": ("size", "مقاس", "المقاس"),
    "qty": ("qty", "quantity", "كمية", "الكمية", "عدد"),
    "width": ("width", "عرض", "العرض"),
    "material": ("material", "خامة", "الخامة"),
    "formats": ("formats", "format", "صيغ", "الصيغ", "صيغة", "export"),
    "time": ("time", "وقت", "الوقت", "seconds"),
    "spacing": ("spacing", "مسافة", "المسافة"),
    "engine": ("engine", "محرك"),
    "buffer_x": ("buffer_x", "بفر طول", "بفر X", "بفر_طول"),
    "buffer_y": ("buffer_y", "بفر عرض", "بفر Y", "بفر_عرض"),
    "shrink_x": ("shrink_x", "انكماش طول", "انكماش_طول"),
    "shrink_y": ("shrink_y", "انكماش عرض", "انكماش_عرض"),
    "attempts": ("attempts", "محاولات", "best_of"),
}
NUMERIC = ("width", "time", "spacing", "buffer_x", "buffer_y", "shrink_x", "shrink_y", "attempts")
FORMATS = ("plt", "dxf", "iso", "scb")


@dataclass
class OrderLine:
    file: str
    size: str
    qty: int


@dataclass
class MarkerOrder:
    name: str
    lines: list[OrderLine] = field(default_factory=list)
    width: float | None = None
    material: str | None = None
    formats: list[str] = field(default_factory=list)
    time: float | None = None
    spacing: float | None = None
    buffer_x: float | None = None
    buffer_y: float | None = None
    shrink_x: float | None = None          # %
    shrink_y: float | None = None
    attempts: float | None = None          # best result: attempts per marker (shortest wins)
    engine: str | None = None
    file_stem: str = ""                    # output file name (unique within the batch)
    # results
    status: str = "pending"            # pending / running / done / failed
    message: str = ""
    length_mm: float | None = None
    efficiency: float | None = None
    outputs: list[str] = field(default_factory=list)


class OrderError(ValueError):
    pass


def _rows(path: str) -> list[list[str]]:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
    return [[c.strip() for c in r] for r in csv.reader(text.splitlines(), dialect)]


def _num(v, kind=float):
    v = (v or "").strip().replace(",", ".")
    if not v:
        return None
    return kind(float(v))


def parse_orders(path: str) -> list[MarkerOrder]:
    rows = [r for r in _rows(path) if any(c for c in r)]
    if not rows:
        raise OrderError("the order list is empty")
    head = [h.strip().lower() for h in rows[0]]
    col = {}
    for key, names in HEADERS.items():
        for i, h in enumerate(head):
            if h in names:
                col[key] = i
                break
    missing = [k for k in ("marker", "file", "size", "qty") if k not in col]
    if missing:
        raise OrderError(f"missing columns: {', '.join(missing)} (first row must contain the column names)")
    base = os.path.dirname(os.path.abspath(path))
    orders: dict[str, MarkerOrder] = {}
    last_marker = last_file = None
    for n, r in enumerate(rows[1:], start=2):
        def get(k):
            i = col.get(k)
            return r[i].strip() if i is not None and i < len(r) else ""
        name = get("marker") or last_marker
        f = get("file") or last_file
        if not name or not f:
            raise OrderError(f"row {n}: marker and file are required")
        last_marker, last_file = name, f
        try:
            qty = _num(get("qty"), int)
        except ValueError:
            raise OrderError(f"row {n}: quantity {get('qty')!r} is not a number")
        size = get("size")
        if not size or not qty:
            continue
        if qty < 0:
            raise OrderError(f"row {n}: quantity {qty} is negative")
        o = orders.setdefault(name, MarkerOrder(name))
        path_f = f if os.path.isabs(f) else os.path.join(base, f)
        o.lines.append(OrderLine(os.path.normpath(path_f), size, qty))
        for k in NUMERIC:
            if getattr(o, k) is None:
                try:
                    setattr(o, k, _num(get(k)))
                except ValueError:
                    raise OrderError(f"row {n}: {k} {get(k)!r} is not a number")
        if o.material is None and get("material"):
            o.material = get("material")
        if not o.formats and get("formats"):
            o.formats = [x for x in re.split(r"[;,\s/+]+", get("formats").lower()) if x]
        if o.engine is None and get("engine"):
            o.engine = get("engine").lower()
    for o in orders.values():
        bad = [x for x in o.formats if x not in FORMATS]
        if bad:
            raise OrderError(f"marker {o.name}: unknown format(s) {', '.join(bad)} (use plt, dxf, iso, scb)")
        if not o.width or o.width <= 0:
            raise OrderError(f"marker {o.name}: fabric width is missing or not positive")
        for k in ("spacing", "buffer_x", "buffer_y"):
            if (getattr(o, k) or 0) < 0:
                raise OrderError(f"marker {o.name}: {k} is negative")
        for k in ("shrink_x", "shrink_y"):
            if abs(getattr(o, k) or 0) > 20:
                raise OrderError(f"marker {o.name}: {k} {getattr(o, k)} % is outside -20 .. +20 %")
    return list(orders.values())


def _match_size(wanted: str, available: list[str]) -> str | None:
    w = wanted.strip().lower()
    for s in available:
        if s.lower() == w:
            return s
    return None


def build_project(order: MarkerOrder, cache: dict | None = None):
    """Session with every design file of the order, quantities set per (file, size)."""
    import copy
    from . import service
    files = list(dict.fromkeys(l.file for l in order.lines))
    if len(files) > service.MAX_MODELS:
        raise OrderError(f"at most {service.MAX_MODELS} design files per marker")
    for f in files:
        if not os.path.isfile(f):
            raise OrderError(f"file not found: {f}")

    def opened(f):
        if cache is not None and f in cache:
            return copy.deepcopy(cache[f])
        p, issues = service.open_any(f)
        if cache is not None:
            cache[f] = copy.deepcopy(p)
        return p
    proj = opened(files[0])
    tags = {}
    if len(files) > 1:
        if not service.model_tags(proj):
            service.namespace_project(proj, service._make_tag(proj.name, []))
        for f in files[1:]:
            other_tag_before = set(service.model_tags(proj))
            service.add_model(proj, f)
            new = [t for t in service.model_tags(proj) if t not in other_tag_before]
            tags[f] = new[0] if new else None
        tags[files[0]] = service.model_tags(proj)[0]
    co = proj.cut_order
    co.size_quantities = {s: 0 for s in proj.size_names}
    for line in order.lines:
        tag = tags.get(line.file) if len(files) > 1 else None
        if tag:
            names = [s.split(":", 1)[1] for s in proj.size_names if s.startswith(tag + ":")]
            s = _match_size(line.size, names)
            key = f"{tag}:{s}" if s else None
        else:
            s = key = _match_size(line.size, proj.size_names)
            names = proj.size_names
        if key is None:
            raise OrderError(f"size {line.size!r} not in {os.path.basename(line.file)} "
                             f"(available: {', '.join(names[:20])})")
        co.size_quantities[key] = co.size_quantities.get(key, 0) + line.qty
    if order.material:
        co.material = order.material if order.material not in ("*", "all", "الكل") else None
    proj.name = order.name
    return proj


def run_order(order: MarkerOrder, outdir: str, plt_config=None, iso_config=None, default_time=120.0,
              progress=None, stop_event=None, cache=None, workers=None, default_attempts=1):
    from . import service
    from .model import FabricSettings, NestSettings
    from .project_storage.store import save_project
    order.status, order.message, order.outputs = "running", "", []
    proj = build_project(order, cache)
    fabric = FabricSettings(width=float(order.width), spacing=float(order.spacing or 0.0),
                            buffer_x=float(order.buffer_x or 0.0), buffer_y=float(order.buffer_y or 0.0),
                            shrink_x=float(order.shrink_x or 0.0), shrink_y=float(order.shrink_y or 0.0))
    st = NestSettings(time_limit=float(order.time or default_time), seed=1, engine=order.engine or "auto",
                      best_of=max(1, int(order.attempts or default_attempts or 1)))
    marker, info = service.run_nesting(proj, fabric, st, progress=progress, stop_event=stop_event, name=order.name,
                                       workers=workers)
    unplaced = [p for p in marker.placements if not p.placed]
    if info.get("stopped"):
        raise OrderError("stopped - marker not exported")
    if info.get("cut_short"):
        raise OrderError(f"time limit too short: {info['cut_short']} piece(s) not placed - give this marker more time")
    if not info["validation"].ok:
        raise OrderError("nested marker failed validation: " + info["validation"].summary())
    if unplaced:                           # never write cutting files with pieces missing
        why = sorted({p.reason for p in unplaced if p.reason})
        raise OrderError(f"{len(unplaced)} piece(s) not placed" + (f" ({'; '.join(why)[:200]})" if why else "")
                         + " - marker not exported; check the fabric width / margins")
    proj.markers.append(marker)
    os.makedirs(outdir, exist_ok=True)
    safe = order.file_stem or _stem(order.name)
    l7n = os.path.join(outdir, safe + ".l7n")
    save_project(proj, l7n)
    order.outputs.append(l7n)
    for fmt in order.formats:
        path = os.path.join(outdir, f"{safe}.{fmt}")
        cfg = plt_config if fmt == "plt" else iso_config if fmt == "iso" else None
        if fmt == "scb":
            if len({l.file for l in order.lines}) > 1:
                raise OrderError("scb: a GRAFIS marker file holds one model; this marker uses several files")
            cfg = service.scb_template(proj)
            if cfg is None:
                raise OrderError("scb: list the model's GRAFIS .scb (not the .mdl) as design file - it is the template")
        service.export_marker(proj, marker, fmt, path, cfg, job_name=safe)
        order.outputs.append(path)
    order.length_mm = marker.stats.get("length_mm")
    order.efficiency = marker.stats.get("efficiency_total_width_pct")
    order.status = "done"
    order.message = f"ok · engine {info.get('engine')}"
    return proj, marker


def write_summary(orders: list[MarkerOrder], path: str):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["marker", "status", "length_mm", "efficiency_%", "message", "files"])
        for o in orders:
            w.writerow([o.name, o.status, f"{o.length_mm:.1f}" if o.length_mm else "",
                        f"{o.efficiency:.2f}" if o.efficiency else "", o.message,
                        " | ".join(os.path.basename(x) for x in o.outputs)])


def _stem(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", name).strip() or "marker"


def _take_turn(slot, stop_event) -> bool:
    """Wait until no other nesting run is active (``slot``); False when stopped while waiting."""
    if slot is None:
        return True
    while not slot.acquire(timeout=0.3):
        if stop_event is not None and stop_event.is_set():
            return False
    return True


def run_batch(orders: list[MarkerOrder], outdir: str, plt_config=None, iso_config=None, default_time=120.0,
              on_order=None, progress=None, stop_event=None, slot=None, default_attempts=1) -> str:
    """Process the orders in sequence; a failing order is reported and the next one continues.
    ``slot``: a semaphore shared with other nesting runs of the application (one run at a time)."""
    cache: dict = {}
    used: set = set()
    for o in orders:                   # 'A/B' and 'A:B' must not overwrite each other's files
        stem, n = _stem(o.name), 2
        while stem.lower() in used:
            stem, n = f"{_stem(o.name)}_{n}", n + 1
        used.add(stem.lower())
        o.file_stem = stem
    for k, o in enumerate(orders):
        if stop_event is not None and stop_event.is_set():
            o.status, o.message = "failed", "stopped"
            continue
        t0 = time.monotonic()
        if on_order:
            on_order(k, o)
        got = _take_turn(slot, stop_event)
        if not got:
            o.status, o.message = "failed", "stopped"
            continue
        try:
            run_order(o, outdir, plt_config, iso_config, default_time, progress, stop_event, cache,
                      workers=slot.workers() if slot is not None and hasattr(slot, "workers") else None,
                      default_attempts=default_attempts)
        except Exception as exc:
            o.status, o.message = "failed", str(exc)
        finally:
            if slot is not None:
                slot.release()
        o.message += f" · {time.monotonic() - t0:.0f} s"
        if on_order:
            on_order(k, o)
    summary = os.path.join(outdir, "batch_summary.csv")
    write_summary(orders, summary)
    return summary


def write_template(path: str):
    rows = [["marker", "file", "size", "qty", "width", "material", "formats", "time", "spacing",
             "buffer_x", "buffer_y", "shrink_x", "shrink_y"],
            ["M001", "pants.mdl", "44", "2", "1500", "O", "plt;iso", "120", "0", "0", "0", "0", "0"],
            ["M001", "pants.mdl", "46", "1", "", "", "", "", "", "", "", "", ""],
            ["M001", "shirt.mdl", "M", "1", "", "", "", "", "", "", "", "", ""],
            ["M002", "shirt.mdl", "L", "3", "1440", "O", "dxf", "60", "", "4", "0", "2", "0"]]
    if path.lower().endswith(".xlsx"):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "markers"
        for r in rows:
            ws.append([int(x) if x.isdigit() else x for x in r])
        wb.save(path)
    else:
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            csv.writer(fh).writerows(rows)
